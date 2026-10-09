import docx
from docx.shared import Pt, Inches, RGBColor

template_path = '02_raw_data/applications/iab_form_en/Antrag_SUF_EN.docx'
output_path = '02_raw_data/applications/Antrag_SUF_EN_Jubilacion_Potsdam.docx'

doc = docx.Document(template_path)

# Dictionary of replacements targeting "Please enter your text here." or placeholder text
replacements = {
    11: """Cristhian David Cáceres Mateus
Academic Researcher / Project Coordinator
Department of Business
University of Europe for Applied Sciences (UE Germany)
Think Campus, Konrad-Zuse-Ring 11, 14469 Potsdam, Germany
E-Mail: cristhian.caceres@ue-germany.de | Phone: +49 157 58042367
(Project Director: Prof. Dr. Talha Ali Khan)""",

    18: "Wealth, Migration & Retirement Sustainability in Germany: Dynamic Microsimulation and Econometric Policy Analysis (2025–2070)",

    20: """1. Problem Statement & Research Objectives:
This research project investigates statutory pension accumulation (GRV, SGB VI) and long-term old-age poverty risks across diverse migration and demographic cohorts in Germany. In the context of demographic aging and shifting labor supply, understanding the pace and quality of labor market integration among immigrant populations is essential. Neuarrivals—particularly from Ukraine (§ 24 AufenthG) and humanitarian source countries—frequently possess substantial formal human capital but initially face deskilling, delayed qualification recognition, and language acquisition hurdles.

This project utilizes microdata from the IAB-BAMF-SOEP Survey of Refugees and the IAB-SOEP Migration Samples to:
(a) Empirically estimate the determinants and timeline of transition into regular social-security-contributing employment (sozialversicherungspflichtige Beschäftigung) over duration of residence.
(b) Quantify the wage penalty of unaccredited foreign qualifications (deskilling) and the premium of CEFR German language acquisition (A1 to C2).
(c) Calibrate a dynamic life-cycle microsimulation platform (2025–2070) that projects annual pension points (Entgeltpunkte), child-rearing credits (§ 56 SGB VI), and old-age social assistance interaction (Grundsicherung im Alter, SGB XII).

2. Core Hypotheses:
- Hypothesis 1 (Assimilation Trajectory): The probability of regular social-security employment increases concavely with residence duration, with formal qualification recognition accelerating wage growth.
- Hypothesis 2 (Deskilling Wage Gap): Foreign tertiary degrees without formal German recognition yield substantial initial gross earnings penalties compared to recognized equivalents.
- Hypothesis 3 (Language Capital): Attaining B2/C-level German proficiency significantly increases monthly gross earnings and annual pension point accumulation.

3. Methodology:
- Econometric Estimation: Extended Mincerian earnings equations and longitudinal panel regressions (Fixed and Random Effects) modeling log gross wages and Entgeltpunkte accumulation with robust standard errors.
- Dynamic Microsimulation: Synthesizing estimated empirical distributions into forward-looking individual lifecycle paths up to the statutory retirement age of 67.""",

    22: """The project directly addresses central questions of scientific labour market research by analyzing:
1. Transition rates and pathways of migrants and refugees into regular, insured employment vs. marginal employment (Minijobs) or unemployment.
2. The returns to human capital, qualification recognition procedures, and host-country language acquisition in the German labor market.
3. The long-term transmission of labor market integration outcomes into statutory pension entitlements and social insurance sustainability.""",

    23: "12/31/2027",

    25: "No. Internal academic research project conducted at the University of Europe for Applied Sciences under faculty supervision (no third-party commercial funding).",

    29: """The requested microdata are strictly necessary because published aggregate statistics lack multivariate cross-tabulations of residence duration, foreign credential recognition status, German language proficiency, and monthly gross wages.

Only the IAB-BAMF-SOEP Survey of Refugees and IAB-SOEP Migration Samples provide:
1. Detailed arrival cohorts, legal residency status, and migration biographies.
2. Verified accreditation procedures for foreign educational and vocational qualifications.
3. Standardized language competence indicators according to the CEFR (A1–C2).
4. Granular monthly gross and net employment earnings necessary to model pension contribution trajectories.

Aggregated data cannot identify the multivariate returns to qualification recognition and language acquisition required to calibrate our microsimulation models.""",

    32: """University of Europe for Applied Sciences GmbH
Potsdam Campus (Think Campus, Konrad-Zuse-Ring 11, 14469 Potsdam, Germany)
Legal form: Gesellschaft mit beschränkter Haftung (GmbH).
Institutional accreditation: State-accredited German university of applied sciences recognized by the Ministry of Science, Research and Culture of Brandenburg (MWFK) and the German Science and Humanities Council (Wissenschaftsrat).""",

    34: """University of Europe for Applied Sciences, Potsdam Campus
Think Campus, Konrad-Zuse-Ring 11, 14469 Potsdam, Germany.
(With secure authorized mobile working from Berlin home office adhering to the university's data security policy and FDZ security guidelines).""",

    37: "State-accredited German higher education institution (University of Europe for Applied Sciences) with full institutional accreditation for independent scientific research and higher education under German federal and state law.",

    39: "None (internal university project).",

    42: """User 1 (Academic Project Director / Supervising Faculty):
Prof. Dr. Talha Ali Khan
Vice President for Research | Professor of Data Science
Department of Business, University of Europe for Applied Sciences
Think Campus, Konrad-Zuse-Ring 11, 14469 Potsdam, Germany
Phone: +49 30 338539710 | E-Mail: talhaali.khan@ue-germany.de

User 2 (Academic Researcher / Project Member):
Cristhian David Cáceres Mateus (Student ID: 93515346)
M.Sc. Data Science / Researcher
University of Europe for Applied Sciences
Ostseestraße 46, 10409 Berlin, Germany
Phone: +49 157 58042367 | E-Mail: cristhian.caceres@ue-germany.de"""
}

for idx, text in replacements.items():
    if idx < len(doc.paragraphs):
        p = doc.paragraphs[idx]
        p.text = text
        # Set clean styling
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)

# Add custom row to Table 1 for IAB-BAMF-SOEP Survey of Refugees and IAB-SOEP Migration Samples
t = doc.tables[0]
new_row = t.add_row()
new_row.cells[0].text = "[X] IAB-BAMF-SOEP Survey of Refugees (Samples M3–M5, Scientific Use File)"
new_row.cells[1].text = "[X] IAB-SOEP Migration Samples (Samples M1–M2, Scientific Use File)"

doc.save(output_path)
print(f"Successfully generated populated application form at: {output_path}")
